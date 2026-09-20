import streamlit as st
import sqlite3
import pandas as pd
from google import genai
from docx import Document
import io
import os
from pypdf import PdfReader
from datetime import datetime, date
import streamlit.components.v1 as components

# ----------------- STREAMLIT PAGE SETUP -----------------
st.set_page_config(page_title="टपाल व्यवस्थापन व कार्यविवरण नोंदवही", layout="wide", page_icon="🏛️")

st.markdown("""
<style>
    .main-header {
        background: linear-gradient(135deg, #1e3c72 0%, #2a5298 100%);
        padding: 18px;
        border-radius: 12px;
        color: white;
        text-align: center;
        margin-bottom: 25px;
    }
</style>
""", unsafe_allow_html=True)

# ----------------- SQLITE DATABASE SETUP -----------------
conn = sqlite3.connect("tapal_data.db", check_same_thread=False)
c = conn.cursor()

c.execute('''
    CREATE TABLE IF NOT EXISTS tapal_entries (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        pra_no TEXT,
        file_no TEXT, 
        subject_code TEXT, 
        inward_no TEXT, 
        computer_no TEXT,
        inward_date TEXT, 
        letter_no_date TEXT, 
        letter_from TEXT, 
        letter_type TEXT,
        subject TEXT, 
        emp_name TEXT, 
        address TEXT, 
        district TEXT,
        action_taken_date TEXT, 
        action_taken TEXT, 
        computer_no_2 TEXT,
        application_close_date TEXT, 
        final_action TEXT, 
        final_action_details TEXT,
        letter_sent_to TEXT, 
        paper_go_in_record TEXT, 
        calling_report TEXT, 
        report_received_date TEXT,
        info_requested_from TEXT,
        office_letter_no_date TEXT,
        reminder_letter_date TEXT,
        remarks TEXT,
        reminder_1 TEXT,
        reminder_2 TEXT,
        reminder_3 TEXT,
        reminder_4 TEXT,
        reminder_discription TEXT
    )
''')
conn.commit()

# Gemini API Key
GEMINI_API_KEY = "YOUR_GEMINI_API_KEY"
client = genai.Client(api_key=GEMINI_API_KEY)

def extract_pdf_text(uploaded_file):
    if uploaded_file is not None:
        reader = PdfReader(uploaded_file)
        text = ""
        for page in reader.pages:
            text += page.extract_text() or ""
        return text
    return ""

def format_to_ddmmyyyy(date_val):
    if pd.isna(date_val) or not date_val or str(date_val).lower() in ['nan', 'nat', 'none', '']:
        return ""
    val_str = str(date_val).strip()
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    if ' ' in val_str:
        val_str = val_str.split(' ')[0]
    
    try:
        dt = datetime.strptime(val_str, '%Y-%m-%d')
        return dt.strftime('%d/%m/%Y')
    except ValueError:
        pass

    try:
        dt = pd.to_datetime(val_str, dayfirst=True, errors='coerce')
        if pd.notna(dt):
            return dt.strftime('%d/%m/%Y')
        return val_str
    except Exception:
        return val_str

def parse_date_safely(date_str):
    if not date_str or str(date_str).lower() in ['nan', 'nat', 'none', '']:
        return None
    val_str = str(date_str).strip()
    if val_str.endswith('.0'):
        val_str = val_str[:-2]
    if ' ' in val_str:
        val_str = val_str.split(' ')[0]
    try:
        dt = pd.to_datetime(val_str, dayfirst=True, errors='coerce')
        return dt.date() if not pd.isna(dt) else None
    except Exception:
        return None

def calculate_pending_days(inward_date_str):
    if not inward_date_str or pd.isna(inward_date_str) or str(inward_date_str).lower() in ['nan', 'nat', 'none', '']:
        return ""
    try:
        inward_dt = pd.to_datetime(inward_date_str, format='%d/%m/%Y', dayfirst=True, errors='coerce')
        if pd.isna(inward_dt):
            inward_dt = pd.to_datetime(inward_date_str, dayfirst=True, errors='coerce')
        if pd.isna(inward_dt):
            return ""
        days = (datetime.now() - inward_dt).days
        return f"{days} दिवस" if days >= 0 else "0 दिवस"
    except Exception:
        return ""

def clean_df(df):
    if df is not None and not df.empty:
        df = df.fillna('')
        df = df.astype(str)
        df = df.replace(['NaN', 'nan', 'NaT', 'NAT', 'None', 'NONE', 'none', '<NA>'], '')
        for col in df.columns:
            df[col] = df[col].apply(lambda x: x[:-2] if x.endswith('.0') else x)
            df[col] = df[col].replace(['nan', 'nat', 'None', ''], '')
    return df

def get_clean_number(val):
    if pd.notna(val):
        val_str = str(val).strip()
        if val_str and val_str.lower() not in ['none', 'nan', 'null', '']:
            if val_str.endswith('.0'):
                val_str = val_str[:-2]
            return val_str
    return ""

def render_html_table(df, title):
    df = clean_df(df)
    table_html = f"""
    <html>
    <head>
        <meta charset="utf-8">
        <style>
            body {{ font-family: Arial, sans-serif; margin: 10px; }}
            h3 {{ text-align: center; margin-bottom: 15px; color: #1e3c72; }}
            table {{ width: 100%; border-collapse: collapse; font-size: 11px; }}
            th, td {{ border: 1px solid #333; padding: 6px 8px; text-align: left; vertical-align: top; word-wrap: break-word; }}
            th {{ background-color: #f0f4f8; text-align: center; font-weight: bold; }}
            @media print {{
                .no-print {{ display: none; }}
                @page {{ size: A4 landscape; margin: 10mm; }}
            }}
            .print-btn {{ background-color: #1E3C72; color: white; padding: 8px 16px; border: none; border-radius: 4px; cursor: pointer; font-size: 14px; margin-bottom: 15px; }}
        </style>
    </head>
    <body>
        <button class="print-btn no-print" onclick="window.print()">🖨️ प्रिंट / PDF म्हणून सेव्ह करा</button>
        <h3>{title}</h3>
        <table>
            <thead>
                <tr>{"".join([f"<th>{col}</th>" for col in df.columns])}</tr>
            </thead>
            <tbody>
    """
    for _, row in df.iterrows():
        table_html += "<tr>"
        for val in row:
            text = str(val) if pd.notna(val) and str(val).strip() != "" and str(val).lower() not in ['none', 'nan', 'nat'] else ""
            table_html += f"<td>{text}</td>"
        table_html += "</tr>"
        
    table_html += "</tbody></table></body></html>"
    components.html(table_html, height=600, scrolling=True)

# ----------------- STREAMLIT HEADER -----------------
st.markdown("""
<div class="main-header">
    <h1>🏛️ महसूल विभाग - टपाल नोंदणी व कार्यविवरण नोंदवही (Worksheet)</h1>
    <p>टपाल नोंद, पुढील कार्यवाही अद्ययावत करणे, गोषवारा व रिपोर्ट्स</p>
</div>
""", unsafe_allow_html=True)

tab1, tab2, tab3, tab4, tab5, tab6, tab7 = st.tabs([
    "📝 नवीन टपाल नोंद", 
    "✏️ कार्यवाही भरणे (Update)",
    "📊 मास्टर रजिस्टर व गोषवारा", 
    "📋 प्रलंबित संदर्भ",
    "⏳ प्रतिक्षाधिन प्रकरणे (Await Reg)",
    "🖨️ कार्यविवरण नोंदवही", 
    "📄 AI टिपणी जनरेटर"
])

# ----------------- TAB 1: ENTRY FORM -----------------
with tab1:
    st.markdown("### 📝 नवीन टपाल नोंदणी फॉर्म")
    with st.form("tapal_form", clear_on_submit=True):
        col1, col2, col3 = st.columns(3)
        
        with col1:
            file_no = st.text_input("File No")
            subject_code = st.text_input("Subject Code")
            inward_no = st.text_input("Inward No (आवक क्रमांक)")
            computer_no = st.text_input("Computer No")
            inward_date_val = st.date_input("Inward Date (प्राप्त दिनांक)", datetime.now())
            inward_date = inward_date_val.strftime("%d/%m/%Y")
            letter_no_date = st.text_input("Letter No & Date")
            letter_from = st.text_input("From (कोणाकडून मिळाले / अर्जदार)")
            
        with col2:
            letter_type = st.selectbox("Letter Type", ["सर्वसाधारण", "न्यायालयीन", "मंत्री", "शासन", "माहिती अधिकार", "इतर"])
            subject = st.text_area("Subject (विषय)")
            emp_name = st.text_input("Emp Name / तक्रारदार")
            address = st.text_input("Address")
            district = st.selectbox("District", ["नाशिक", "धुळे", "जळगाव", "नंदुरबार", "अहिल्यानगर", "इतर"])
            
        with col3:
            computer_no_2 = st.text_input("Computer No 2 (इतर संदर्भ)")
            remarks = st.text_area("शेरा (Remarks)")

        submit_btn = st.form_submit_button("💾 टपाल नोंद सेव्ह करा", type="primary")

        if submit_btn:
            clean_inw = str(inward_no).strip()
            clean_comp = str(computer_no).strip()
            
            c.execute('''
                INSERT INTO tapal_entries (
                    file_no, subject_code, inward_no, computer_no, inward_date,
                    letter_no_date, letter_from, letter_type, subject, emp_name,
                    address, district, action_taken_date, action_taken, computer_no_2,
                    application_close_date, final_action, final_action_details, letter_sent_to,
                    paper_go_in_record, calling_report, report_received_date,
                    info_requested_from, office_letter_no_date, reminder_letter_date, remarks
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                str(file_no), str(subject_code), clean_inw, clean_comp, str(inward_date),
                str(letter_no_date), str(letter_from), str(letter_type), str(subject), str(emp_name),
                str(address), str(district), "", "", str(computer_no_2),
                "", "", "", "", "नाही", "नाही", "", "", "", "", str(remarks)
            ))
            conn.commit()
            st.success("✅ टपाल नोंद यशस्वीरीत्या सेव्ह झाली आहे!")

# ----------------- TAB 2: UPDATE ACTION -----------------
with tab2:
    st.markdown("### ✏️ प्रकरणाची पुढील कार्यवाही नोंदवणे / अपडेट करणे")
    
    df_all_tapal = pd.read_sql_query("SELECT id, inward_no, subject, inward_date, action_taken FROM tapal_entries ORDER BY id DESC", conn)
    df_all_tapal = clean_df(df_all_tapal)
    
    if not df_all_tapal.empty:
        tapal_options = {
            f"ID {row['id']} | आवक क्र: {get_clean_number(row['inward_no']) or 'N/A'} | विषय: {str(row['subject'])[:30]}...": row['id']
            for _, row in df_all_tapal.iterrows()
        }
        
        selected_option = st.selectbox("🎯 कार्यवाही अपडेट करण्यासाठी प्रकरण निवडा:", list(tapal_options.keys()))
        selected_id = tapal_options[selected_option]
        
        entry = pd.read_sql_query("SELECT * FROM tapal_entries WHERE id = ?", conn, params=(selected_id,)).iloc[0]
        
        st.markdown("---")
        st.subheader("📌 प्रकरणाची सद्यस्थिती व माहिती")
        
        col_info1, col_info2, col_info3 = st.columns(3)
        col_info1.write(f"**आवक क्रमांक:** {get_clean_number(entry['inward_no']) or '-'}")
        col_info1.write(f"**आवक दिनांक:** {format_to_ddmmyyyy(entry['inward_date']) or '-'}")
        col_info2.write(f"**अर्जदार/प्राधिकरण:** {entry['emp_name'] or entry['letter_from'] or '-'}")
        col_info2.write(f"**पत्राचा प्रकार:** {entry['letter_type'] or '-'}")
        col_info3.write(f"**विषय:** {entry['subject'] or '-'}")
        
        st.markdown("---")
        st.markdown("#### 🔄 नवीन कार्यवाही व अंतिम तपशील अपडेट करा:")
        
        with st.form("update_action_form"):
            col_u1, col_u2, col_u3 = st.columns(3)
            
            with col_u1:
                existing_act_date = parse_date_safely(entry['action_taken_date'])
                u_act_date_val = st.date_input("Action Taken Date (कार्यवाही दिनांक):", value=existing_act_date)
                u_action_taken_date = u_act_date_val.strftime("%d/%m/%Y") if u_act_date_val else ""
                
                raw_act_taken = entry['action_taken'] if pd.notna(entry['action_taken']) and str(entry['action_taken']).lower() not in ['nan', 'nat', 'none'] else ""
                u_action_taken = st.text_input("केलेली कार्यवाही (Action Taken - भरल्यास प्रकरण निपटारा होईल):", value=raw_act_taken)
                
                raw_info_req = entry['info_requested_from'] if pd.notna(entry['info_requested_from']) and str(entry['info_requested_from']).lower() not in ['nan', 'nat', 'none'] else ""
                u_info_requested_from = st.text_input("कोणाकडून माहिती मागविली:", value=raw_info_req)
                
                raw_final_act = entry['final_action'] if pd.notna(entry['final_action']) and str(entry['final_action']).lower() not in ['nan', 'nat', 'none'] else ""
                u_final_action = st.text_input("अंतिम कार्यवाही प्रकार (Final Action):", value=raw_final_act)

            with col_u2:
                raw_off_let = entry['office_letter_no_date'] if pd.notna(entry['office_letter_no_date']) and str(entry['office_letter_no_date']).lower() not in ['nan', 'nat', 'none'] else ""
                u_office_letter_no_date = st.text_input("कार्यालयीन पत्र क्र. व दिनांक:", value=raw_off_let)
                
                existing_rem_date = parse_date_safely(entry['reminder_letter_date'])
                u_rem_date_val = st.date_input("तगादा केला असेल तर पत्र दिनांक:", value=existing_rem_date)
                u_reminder_letter_date = u_rem_date_val.strftime("%d/%m/%Y") if u_rem_date_val else ""
                
                raw_let_sent = entry['letter_sent_to'] if pd.notna(entry['letter_sent_to']) and str(entry['letter_sent_to']).lower() not in ['nan', 'nat', 'none'] else ""
                u_letter_sent_to = st.text_input("पत्र कोणाला पाठवले (Sent To):", value=raw_let_sent)
                
                raw_close_date = entry['application_close_date'] if pd.notna(entry['application_close_date']) and str(entry['application_close_date']).lower() not in ['nan', 'nat', 'none'] else ""
                u_application_close_date = st.text_input("प्रकरण बंद दिनांक (Close Date):", value=raw_close_date)

            with col_u3:
                paper_rec_bool = True if str(entry['paper_go_in_record']).strip() == 'होय' else False
                u_paper_rec = st.checkbox("Paper Go in Record (अभिलेखात जमा)", value=paper_rec_bool)
                
                call_rep_bool = True if str(entry['calling_report']).strip() == 'होय' else False
                u_call_rep = st.checkbox("Calling Report (अहवाल मागविला)", value=call_rep_bool)
                
                raw_final_det = entry['final_action_details'] if pd.notna(entry['final_action_details']) and str(entry['final_action_details']).lower() not in ['nan', 'nat', 'none'] else ""
                u_final_action_details = st.text_area("अंतिम सविस्तर शेरा (Final Details):", value=raw_final_det)
                
                raw_remarks = entry['remarks'] if pd.notna(entry['remarks']) and str(entry['remarks']).lower() not in ['nan', 'nat', 'none'] else ""
                u_remarks = st.text_area("सामान्य शेरा (Remarks):", value=raw_remarks)

            update_btn = st.form_submit_button("💾 अद्ययावत करा (Save Changes)", type="primary")

            if update_btn:
                paper_go_val = "होय" if u_paper_rec else "नाही"
                calling_rep_val = "होय" if u_call_rep else "नाही"
                
                c.execute('''
                    UPDATE tapal_entries SET
                        action_taken_date = ?,
                        action_taken = ?,
                        info_requested_from = ?,
                        office_letter_no_date = ?,
                        reminder_letter_date = ?,
                        final_action = ?,
                        final_action_details = ?,
                        letter_sent_to = ?,
                        application_close_date = ?,
                        paper_go_in_record = ?,
                        calling_report = ?,
                        remarks = ?
                    WHERE id = ?
                ''', (
                    u_action_taken_date, u_action_taken, u_info_requested_from,
                    u_office_letter_no_date, u_reminder_letter_date, u_final_action,
                    u_final_action_details, u_letter_sent_to, u_application_close_date,
                    paper_go_val, calling_rep_val, u_remarks, selected_id
                ))
                conn.commit()
                st.success("✅ टपालाची पुढील कार्यवाही यशस्वीरीत्या अपडेट झाली आहे!")
                st.rerun()
    else:
        st.info("डेटाबेसमध्ये अपडेट करण्यासाठी कोणत्याही टपाल नोंदी नाहीत.")

# ----------------- TAB 3: REGISTER & GOSHWARA -----------------
with tab3:
    st.markdown("### 📈 टपाल गोषवारा व डेटा व्यवस्थापन")
    df_all = pd.read_sql_query("SELECT * FROM tapal_entries", conn)
    df_all = clean_df(df_all)
    
    if not df_all.empty:
        pending_count = len(df_all[df_all['action_taken'].astype(str).str.strip() == ''])
        completed_count = len(df_all) - pending_count
        
        m1, m2, m3, m4, m5, m6 = st.columns(6)
        m1.metric("एकूण नोंद टपाल", len(df_all))
        m2.metric("⚠️ प्रलंबित संदर्भ", pending_count)
        m3.metric("✅ निपटारा झालेले", completed_count)
        m4.metric("शासन संदर्भ", len(df_all[df_all['letter_type'] == 'शासन']))
        m5.metric("मंत्री संदर्भ", len(df_all[df_all['letter_type'] == 'मंत्री']))
        m6.metric("न्यायालयीन संदर्भ", len(df_all[df_all['letter_type'] == 'न्यायालयीन']))

        st.markdown("---")
        col_acc1, col_acc2 = st.columns(2)
        with col_acc1:
            st.subheader("📌 पत्राच्या प्रकारानुसार गोषवारा")
            summary_type = df_all.groupby('letter_type').size().reset_index(name='संख्या (Count)')
            st.dataframe(summary_type, use_container_width=True)
            
        with col_acc2:
            st.subheader("📍 प्रलंबिततेनुसार गोषवारा")
            st.write(f"- **प्रलंबित संदर्भ:** {pending_count}")
            st.write(f"- **निपटारा झालेले:** {completed_count}")

    st.markdown("---")
    st.markdown("### 📥📤 Data Import/ Export")
    
    col_dl, col_ul = st.columns(2)
    
    with col_dl:
        st.subheader("डेटा डाउनलोड करा")
        if not df_all.empty:
            csv_data = df_all.to_csv(index=False).encode('utf-8-sig')
            st.download_button(
                label="📥 सर्व डेटा Excel (CSV) डाउनलोड करा",
                data=csv_data,
                file_name=f"Tapal_Master_Register_{datetime.now().strftime('%Y%m%d')}.csv",
                mime="text/csv",
                type="primary"
            )
        else:
            st.info("डाउनलोड करण्यासाठी डेटा उपलब्ध नाही.")
            
    with col_ul:
        st.subheader("नवीन Excel किंवा CSV फाईल अपलोड करा")
        uploaded_file = st.file_uploader("तुमची नवीन Excel किंवा CSV फाईल निवडा", type=["csv", "xlsx", "xls"], key="new_file_uploader")
        
        replace_existing = st.checkbox("⚠️ जुना सर्व डेटा डिलीट करून नवीन फाईलचा डेटा टाका (Overwrite)", value=True)
        
        if uploaded_file is not None:
            if st.button("🚀 फाईल अपलोड व अपडेट करा", type="primary"):
                try:
                    if uploaded_file.name.endswith('.csv'):
                        df_upload = pd.read_csv(uploaded_file)
                    else:
                        df_upload = pd.read_excel(uploaded_file)
                        
                    # कॉलमचे स्पेलिंग आणि स्पेसेस पूर्णपणे साफ करणे
                    df_upload.columns = [str(c).strip() for c in df_upload.columns]
                    
                    # अचूक आणि स्पष्ट कॉलम मॅपिंग (Subject Code आणि Subject वेगळे राहण्यासाठी)
                    rename_map = {
                        'ID': 'id_excel',
                        'Pra No': 'pra_no',
                        'Subject Code': 'subject_code',
                        'Inward No': 'inward_no',
                        'Inward Date': 'inward_date',
                        'Letter No & Date': 'letter_no_date',
                        'From': 'letter_from',
                        'Letter Type': 'letter_type',
                        'Subject': 'subject',
                        'Emp Name': 'emp_name',
                        'Address': 'address',
                        'District': 'district',
                        'Action Taken Date': 'action_taken_date',
                        'Action Taken': 'action_taken',
                        'Application Close Date': 'application_close_date',
                        'Final Action': 'final_action',
                        'Letter Sent to': 'letter_sent_to',
                        'Paper Go to Record': 'paper_go_in_record',
                        'Calling Report': 'calling_report',
                        'Report Received Date': 'report_received_date',
                        'Reminder 1': 'reminder_1',
                        'Reminder 2': 'reminder_2',
                        'Reminder 3': 'reminder_3',
                        'Reminder 4': 'reminder_4',
                        'Reminder Discription': 'reminder_discription'
                    }
                    
                    df_upload = df_upload.rename(columns=rename_map)

                    # अतिरिक्त खबरदारीसाठी स्मार्ट मॅचिंग (Subject Code आणि Subject ची गफलत टाळण्यासाठी)
                    for col in list(df_upload.columns):
                        col_lower = str(col).lower()
                        if col_lower in ['subject code', 'subject_code', 'subjectcode']:
                            df_upload = df_upload.rename(columns={col: 'subject_code'})
                        elif col_lower in ['subject'] or (('subject' in col_lower or 'विषय' in col) and 'code' not in col_lower):
                            df_upload = df_upload.rename(columns={col: 'subject'})
                        elif 'inward' in col_lower and 'no' in col_lower:
                            df_upload = df_upload.rename(columns={col: 'inward_no'})
                        elif 'inward' in col_lower and 'date' in col_lower:
                            df_upload = df_upload.rename(columns={col: 'inward_date'})
                        elif 'action' in col_lower and 'date' in col_lower:
                            df_upload = df_upload.rename(columns={col: 'action_taken_date'})
                        elif 'action' in col_lower and 'date' not in col_lower:
                            df_upload = df_upload.rename(columns={col: 'action_taken'})

                    if 'id' in df_upload.columns:
                        df_upload = df_upload.drop(columns=['id'])
                    if 'id_excel' in df_upload.columns:
                        df_upload = df_upload.drop(columns=['id_excel'])

                    cursor = conn.cursor()
                    cursor.execute("PRAGMA table_info(tapal_entries)")
                    db_columns = [row[1] for row in cursor.fetchall() if row[1] != 'id']

                    for col in db_columns:
                        if col not in df_upload.columns:
                            df_upload[col] = ""
                            
                    df_upload = df_upload[[col for col in db_columns if col in df_upload.columns]]
                    
                    if replace_existing:
                        cursor.execute("DELETE FROM tapal_entries")
                        conn.commit()
                    
                    df_upload.to_sql('tapal_entries', conn, if_exists='append', index=False)
                    st.success("✅ Subject Code आणि Subject सह सर्व डेटा अचूक जागेवर सेव्ह झाला आहे!")
                    st.rerun()
                except Exception as e:
                    st.error(f"फाइल अपलोड करताना त्रुटी आली: {e}")

    st.markdown("---")
    st.markdown("### 📋 टपाल मास्टर रजिस्टर")
    df = pd.read_sql_query("SELECT * FROM tapal_entries ORDER BY id DESC", conn)
    df = clean_df(df)
    if 'inward_no' in df.columns:
        df['inward_no'] = df['inward_no'].apply(get_clean_number)
    if 'computer_no' in df.columns:
        df['computer_no'] = df['computer_no'].apply(get_clean_number)
    if 'inward_date' in df.columns:
        df['inward_date'] = df['inward_date'].apply(format_to_ddmmyyyy)
    if 'action_taken_date' in df.columns:
        df['action_taken_date'] = df['action_taken_date'].apply(format_to_ddmmyyyy)
    
    st.dataframe(df, use_container_width=True)

# ----------------- TAB 4: PENDING REFERENCES -----------------
with tab4:
    st.markdown("### 📋 प्रलंबित संदर्भांची यादी")
    
    df_p = pd.read_sql_query("""
        SELECT * FROM tapal_entries 
        WHERE action_taken IS NULL OR action_taken = '' OR action_taken = 'None' OR action_taken = 'nan'
        ORDER BY id ASC
    """, conn)
    df_p = clean_df(df_p)
    
    if not df_p.empty:
        df_p = df_p[df_p['action_taken'].astype(str).str.strip() == '']
    
    pending_count_tab4 = len(df_p)
    st.metric("⚠️ एकूण प्रलंबित संदर्भ", pending_count_tab4)
    st.markdown("---")
    
    if not df_p.empty:
        df_p.reset_index(inplace=True, drop=True)
        df_p['अ.क्र.'] = df_p.index + 1
        
        df_p['आवक क्रमांक'] = df_p['inward_no'].apply(get_clean_number)
        df_p['inward_date'] = df_p['inward_date'].apply(format_to_ddmmyyyy)
        df_p['प्रलंबित कालावधी'] = df_p['inward_date'].apply(calculate_pending_days)
        
        df_p_display = df_p.rename(columns={
            'subject': 'विषय',
            'inward_date': 'दाखल दिनांक',
            'letter_from': 'कोणाकडून मिळाले',
            'letter_type': 'पत्राचा प्रकार'
        })[['अ.क्र.', 'आवक क्रमांक', 'कोणाकडून मिळाले', 'पत्राचा प्रकार', 'विषय', 'दाखल दिनांक', 'प्रलंबित कालावधी']]

        df_p_display = clean_df(df_p_display)
        st.dataframe(df_p_display, use_container_width=True)
        
        st.markdown("---")
        if st.checkbox("🖨️ प्रलंबित संदर्भ प्रिंट / PDF व्ह्यू दाखवा", key="print_pending"):
            render_html_table(df_p_display, "प्रलंबित संदर्भांची यादी")
    else:
        st.info("सध्या Action Taken रिक्त असलेली कोणतीही प्रलंबित प्रकरणे नाहीत.")

# ----------------- TAB 5: AWAIT REGISTER -----------------
with tab5:
    st.markdown("### ⏳ प्रतिक्षाधिन प्रकरणांची नोंदवही (महसूल आस्थापना) संकलन- (आस्था-५)")
    
    df_await = pd.read_sql_query("""
        SELECT letter_from, inward_date, subject, info_requested_from, 
               office_letter_no_date, reminder_letter_date, remarks, action_taken_date, action_taken 
        FROM tapal_entries 
        WHERE calling_report = 'होय' OR (info_requested_from IS NOT NULL AND info_requested_from != '')
        ORDER BY id ASC
    """, conn)
    df_await = clean_df(df_await)
    
    if not df_await.empty:
        df_await.reset_index(inplace=True, drop=True)
        df_await['अ.क्र.'] = df_await.index + 1
        
        df_await['तक्रारदाराचे नाव / प्राधिकरणाचे नाव / प्राप्त दिनांक'] = df_await.apply(
            lambda r: f"{r['letter_from'] or ''}, दि. {format_to_ddmmyyyy(r['inward_date'])}".strip(), axis=1
        )

        df_await['कार्यालयीन पत्र क्र. व दिनांक'] = df_await.apply(
            lambda r: f"दि. {format_to_ddmmyyyy(r['action_taken_date'])} रोजी {r['action_taken'] or ''}".strip() if r['action_taken_date'] and r['action_taken'] else (f"दि. {format_to_ddmmyyyy(r['action_taken_date'])}" if r['action_taken_date'] else (r['action_taken'] or '')), axis=1
        )

        df_await_display = df_await.rename(columns={
            'subject': 'विषय',
            'info_requested_from': 'कोणाकडून माहिती मागविली',
            'reminder_letter_date': 'तगादा केला असेल तर पत्र दिनांक',
            'remarks': 'शेरा'
        })[['अ.क्र.', 'तक्रारदाराचे नाव / प्राधिकरणाचे नाव / प्राप्त दिनांक', 'विषय', 'कोणाकडून माहिती मागविली', 'कार्यालयीन पत्र क्र. व दिनांक', 'तगादा केला असेल तर पत्र दिनांक', 'शेरा']]

        df_await_display = clean_df(df_await_display)
        st.dataframe(df_await_display, use_container_width=True)
        
        st.markdown("---")
        if st.checkbox("🖨️ प्रतिक्षाधिन प्रकरणे प्रिंट / PDF व्ह्यू दाखवा", key="print_await"):
            render_html_table(df_await_display, "प्रतिक्षाधिन प्रकरणांची नोंदवही (महसूल आस्थापना) संकलन- (आस्था-५)")
    else:
        st.info("सध्या कोणतीही प्रतिक्षाधिन प्रकरणे नोंदवलेली नाहीत.")

# ----------------- TAB 6: WORKSHEET GENERATOR -----------------
with tab6:
    st.markdown("### 🖨️ कार्यविवरण नोंदवही (महसूल आस्थापना) संकलन - (आस्था-५) Worksheet")
    filter_type = st.selectbox("प्रिंटसाठी संदर्भ प्रकार निवडा:", ["सर्व टपाल", "शासन", "मंत्री", "न्यायालयीन", "माहिती अधिकार", "सर्वसाधारण"])
    
    ws_query = "SELECT inward_no, inward_date, letter_no_date, letter_from, letter_type, subject, action_taken_date, action_taken FROM tapal_entries"
    if filter_type != "सर्व टपाल":
        ws_query += f" WHERE letter_type = '{filter_type}'"
    ws_query += " ORDER BY id ASC"
    
    df_ws_raw = pd.read_sql_query(ws_query, conn)
    df_ws_raw = clean_df(df_ws_raw)
    
    if not df_ws_raw.empty:
        df_ws_raw.reset_index(inplace=True, drop=True)
        df_ws_raw['अ.क्र.'] = df_ws_raw.index + 1
        
        if 'inward_no' in df_ws_raw.columns:
            df_ws_raw['inward_no'] = df_ws_raw['inward_no'].apply(get_clean_number)
        if 'inward_date' in df_ws_raw.columns:
            df_ws_raw['inward_date'] = df_ws_raw['inward_date'].apply(format_to_ddmmyyyy)
        if 'action_taken_date' in df_ws_raw.columns:
            df_ws_raw['action_taken_date'] = df_ws_raw['action_taken_date'].apply(format_to_ddmmyyyy)

        df_ws = df_ws_raw.rename(columns={
            'inward_no': 'आवक क्रमांक', 'inward_date': 'आवक दिनांक',
            'letter_no_date': 'पत्र क्र. व दिनांक', 'letter_from': 'कोणाकडून मिळाले',
            'letter_type': 'पत्राचा प्रकार', 'subject': 'पत्राचा विषय',
            'action_taken_date': 'कार्यवाही दिनांक', 'action_taken': 'केलेली कार्यवाही'
        })[['अ.क्र.', 'आवक क्रमांक', 'आवक दिनांक', 'पत्र क्र. व दिनांक', 'कोणाकडून मिळाले', 'पत्राचा प्रकार', 'पत्राचा विषय', 'कार्यवाही दिनांक', 'केलेली कार्यवाही']]
        
        df_ws = clean_df(df_ws)
        st.dataframe(df_ws, use_container_width=True)
        if st.checkbox("🖨️ कार्यविवरण नोंदवही प्रिंट / PDF व्ह्यू दाखवा", key="print_worksheet"):
            render_html_table(df_ws, f"कार्यविवरण नोंदवही संकलन- (आस्था-५) - {filter_type}")

# ----------------- TAB 7: AI GENERATOR -----------------
with tab7:
    st.markdown("### 📄 ऑटोमॅटिक टिपणी व पत्र मसुदा जनरेटर")
    ai_subject = st.text_input("१. कामाचा / पत्राचा विषय:")
    previous_pdf = st.file_uploader("𝟐. फाईलचा पूर्व इतिहास / जुन्या टिपण्या (PDF):", type=["pdf"])
    new_letter_pdf = st.file_uploader("३. नवीन आलेले पत्र (PDF):", type=["pdf"])
    additional_remarks = st.text_area("४. अतिरिक्त निर्देश / शेरा (ऐच्छिक):")
    
    if st.button("🚀 टिपणी व पत्र तयार करा", type="primary"):
        if not ai_subject or not new_letter_pdf:
            st.error("कृपया विषय आणि नवीन आलेले पत्र अपलोड करा.")
        else:
            with st.spinner("Gemini AI टिपणी तयार करत आहे..."):
                try:
                    old_history_text = extract_pdf_text(previous_pdf)
                    new_letter_text = extract_pdf_text(new_letter_pdf)

                    prompt_content = f"तुम्ही एक वरिष्ठ शासकीय अधिकारी/लिपीक आहात. खालील माहितीच्या आधारे शुद्ध, अधिकृत मराठी शासकीय भाषेत टिपणी (Note Sheet) आणि उत्तराचा मसुदा तयार करा.\n\nविषय: {ai_subject}\nमागील इतिहास/टिपणी: {old_history_text}\nनवीन आलेले पत्र: {new_letter_text}\nअतिरिक्त शेरा: {additional_remarks}\n\nखालील फॉरमॅटमध्ये उत्तर द्या:\n---कार्यालयीन टिपणी---\n[सविस्तर टिपणी]\n\n---जावक पत्र मसुदा---\n[जावक पत्राचा मसुदा]"

                    response = client.models.generate_content(
                        model="gemini-2.5-flash",
                        contents=prompt_content,
                    )

                    ai_output = response.text
                    doc = Document()
                    doc.add_heading("कार्यालयीन टिपणी व जावक पत्र", level=1)
                    doc.add_paragraph(ai_output)

                    doc_io = io.BytesIO()
                    doc.save(doc_io)

                    st.success("✅ टिपणी आणि पत्र यशस्वीरीत्या तयार झाले आहे!")
                    st.download_button(
                        label="📥 Word फाईल डाउनलोड करा",
                        data=doc_io.getvalue(),
                        file_name=f"Note_Sheet_{ai_subject[:10]}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                    st.markdown(ai_output)
                except Exception as err:
                    st.error(f"अडचण आली आहे: {err}")
